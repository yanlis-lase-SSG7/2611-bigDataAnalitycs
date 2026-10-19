"""Check the local Word draft against measured Assignment II results."""

import csv
import json
from pathlib import Path
import unittest

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "Big Data Analytics - Laporan 2702751284.docx"
REPORT = ROOT / "Big Data Analytics - Laporan Assignment II 2702751284.docx"
OUT = ROOT / "outputs_assignment2"


@unittest.skipUnless(REPORT.exists() and SOURCE.exists() and
                     (OUT / "model_comparison.csv").exists(),
                     "Local Word manuscript or measured results unavailable")
class Assignment2ReportTests(unittest.TestCase):
    def test_six_report_rows_match_saved_csv_and_validation_json(self):
        document = Document(REPORT)
        tables = [table for table in document.tables if
                  [cell.text for cell in table.rows[0].cells[:3]] ==
                  ["Split", "Model", "F1 validasi"]]
        self.assertEqual(len(tables), 1)
        rows = tables[0].rows[1:]
        self.assertEqual(len(rows), 6)
        label = {"spark_random_forest": "Spark RF",
                 "sklearn_random_forest": "sklearn RF",
                 "sklearn_logistic_regression": "sklearn LR"}
        report = {(r.cells[0].text, r.cells[1].text):
                  [cell.text for cell in r.cells[2:]] for r in rows}
        with (OUT / "model_comparison.csv").open(newline="", encoding="utf-8") as file:
            measured = list(csv.DictReader(file))
        self.assertEqual(len(measured), 6)
        for row in measured:
            key = (row["split"], label[row["model"]])
            with self.subTest(key=key):
                saved = json.loads((OUT / f"{row['split']}_{row['model']}_metrics.json")
                                   .read_text(encoding="utf-8"))
                expected = [saved["validation_selection"]["f1"],
                            *[float(row[field]) for field in (
                                "precision", "recall", "f1", "pr_auc_trapezoid", "roc_auc")]]
                self.assertEqual(report[key], [f"{number:.4f}" for number in expected])

    def test_baseline_is_labeled_and_student_specified_form_is_updated(self):
        source = Document(SOURCE)
        report = Document(REPORT)
        body = "\n".join(p.text for p in report.paragraphs)
        self.assertIn("Gambar 2.2 Lima feature importance terbesar pada baseline Spark RF Assignment I", body)
        self.assertIn("Gambar 6.2 Volume dan delay rate relasi asal–tujuan", body)
        self.assertNotIn("Laporan Assignment I ini membahas", body)
        self.assertIn("pengujian temporal Assignment II dilaporkan pada Bab 5", body)
        self.assertIn("Jakarta, 16 Oktober 2026", body)
        self.assertNotIn("Jakarta, 28 September 2026", body)
        self.assertTrue(source.tables[-2].cell(4, 1).text.startswith(": PENENTUAN STUDI KASUS"))
        header, activity = report.tables[-2:]
        self.assertEqual(header.cell(4, 1).text,
                         ": Assignment II - Pengolahan Big Data Menggunakan Machine Learning dan Graph Analytics")
        self.assertEqual(header.cell(5, 1).text, "Type 2 / Partial Use")
        self.assertEqual([activity.cell(i, 1).text for i in range(1, 6)], ["Yes"] * 5)
        self.assertIn("PySpark MLlib", activity.cell(2, 3).text)
        self.assertIn("NetworkX", activity.cell(2, 3).text)
        self.assertIn("300 DPI", activity.cell(3, 3).text)
        self.assertIn("tabel hasil ML", activity.cell(5, 3).text)
        self.assertIn("meninjau interpretasi akhir", activity.cell(2, 3).text)
        self.assertEqual(activity.cell(6, 0).text,
                         "Approximate contribution of AI in the final submission of assignment is 15%")
        form_start = next(i for i, p in enumerate(report.paragraphs)
                          if "ARTIFICIAL INTELLIGENCE USE DECLARATION FORM" in p.text)
        self.assertEqual(sum(len(p._p.xpath(".//w:drawing"))
                             for p in report.paragraphs[form_start + 1:]), 0)


if __name__ == "__main__":
    unittest.main()
